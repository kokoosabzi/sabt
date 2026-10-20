from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook
from io import BytesIO

from openpyxl import load_workbook

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.auth.dependencies import current_user, get_db
from app.db.base import Base
from app.db.models import (Operation, OperationType, Permission, Project, ProjectAccess, Role, RolePermission, User, UserRole, WorkflowEvent)
from app.main import app


def _context():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = Session(engine)
    now = datetime.now(timezone.utc)
    user = User(username="report-user", password_hash="unused", display_name="Report User",
                active=True, created_at=now, updated_at=now, version=1)
    project = Project(code="REPORT-P", contract_code="77", name="Report Project", status="ACTIVE",
                      timezone="Asia/Tehran", created_at=now, updated_at=now, version=1)
    role = Role(code="report-viewer", name="Report Viewer", active=True)
    permissions = [
        Permission(code="management_dashboard.view", name="Management Dashboard"),
        Permission(code="workflow_report.view", name="Workflow Report"),
    ]
    db.add_all([user, project, role, *permissions])
    db.flush()
    db.add_all([
        RolePermission(role_id=role.id, permission_id=p.id) for p in permissions
    ])
    db.add(UserRole(user_id=user.id, role_id=role.id, project_id=None))
    db.add(ProjectAccess(user_id=user.id, project_id=project.id, active=True,
                         granted_at=now, version=1))
    db.commit()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[current_user] = lambda: user
    return engine, db, TestClient(app), project


def test_management_summary_api_returns_verified_empty_metrics():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/management/summary", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 200
        body = response.json()
        assert body["project"]["id"] == project.id
        assert body["indicators"]["created_operations"] == 0
        assert body["indicators"]["finalized_operations"] == 0
        assert body["period"]["range_semantics"] == "[start_at, end_at)"
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_workflow_report_api_returns_paginated_empty_rows():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/workflow", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
            "limit": 10,
            "offset": 0,
        })
        assert response.status_code == 200
        body = response.json()
        assert body["pagination"] == {"limit": 10, "offset": 0, "total": 0}
        assert body["items"] == []
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_reporting_api_rejects_invalid_range():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/management/summary", params={
            "project_id": project.id,
            "start_at": "2026-11-01T00:00:00+00:00",
            "end_at": "2026-10-01T00:00:00+00:00",
        })
        assert response.status_code == 422
        assert response.json()["detail"] == "report_range_invalid"
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_reporting_api_enforces_project_access():
    engine, db, client, project = _context()
    try:
        db.query(ProjectAccess).delete()
        db.commit()
        response = client.get("/reports/management/summary", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 403
        assert response.json()["detail"] == "forbidden"
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()



def test_workflow_events_api_returns_persisted_events_only():
    engine, db, client, project = _context()
    try:
        now = datetime.now(timezone.utc)
        operation_type = OperationType(
            code="EVENT-TEST", name="Event Test", active=True, system_defined=False,
            created_at=now, updated_at=now,
        )
        db.add(operation_type)
        db.flush()
        operation = Operation(
            project_id=project.id, operation_type_id=operation_type.id,
            status="DRAFT", current_workflow_state="DRAFT", created_by=1,
            created_at=now, updated_at=now, version=1, current_version_number=0,
        )
        db.add(operation)
        db.flush()
        event = WorkflowEvent(
            project_id=project.id, operation_id=operation.id, actor_user_id=1,
            event_type="OPERATION_CREATED", from_status=None, to_status="DRAFT",
            occurred_at=now, metadata_json=None,
        )
        db.add(event)
        db.commit()

        response = client.get("/reports/workflow/events", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
            "event_type": "OPERATION_CREATED",
        })
        assert response.status_code == 200
        body = response.json()
        assert body["pagination"]["total"] == 1
        assert body["items"][0]["event_type"] == "OPERATION_CREATED"
        assert body["items"][0]["to_status"] == "DRAFT"
        assert body["period"]["basis"] == "workflow_events.occurred_at"
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_workflow_csv_export_is_utf8_and_permission_scoped():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/workflow/export.csv", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 200
        assert response.content.startswith(b"\xef\xbb\xbf")
        assert "attachment; filename=\"sabt-workflow-report.csv\"" == response.headers["content-disposition"]
        assert "شناسه عملیات".encode("utf-8") in response.content
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_workflow_events_csv_export_is_available_and_project_scoped():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/workflow/events/export.csv", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 200
        assert response.content.startswith(b"\xef\xbb\xbf")
        assert "sabt-workflow-events.csv" in response.headers["content-disposition"]
        assert "شناسه رخداد".encode("utf-8") in response.content
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_workflow_csv_export_enforces_project_permission():
    engine, db, client, project = _context()
    try:
        db.query(ProjectAccess).delete()
        db.commit()
        response = client.get("/reports/workflow/export.csv", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 403
        assert response.json()["detail"] == "forbidden"
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_csv_export_neutralizes_spreadsheet_formula_cells():
    from app.reports.routes import _csv_response

    response = _csv_response("test.csv", [("value", "value")], [
        {"value": "=HYPERLINK(\"https://example.invalid\",\"click\")"},
        {"value": "ordinary text"},
        {"value": 42},
    ])
    body = response.body.decode("utf-8-sig")
    assert "'=HYPERLINK(" in body
    assert "ordinary text" in body
    assert "42" in body


def test_report_project_context_returns_authorized_project_timezone():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/project-context", params={"project_id": project.id})
        assert response.status_code == 200
        assert response.json()["project"]["timezone"] == "Asia/Tehran"
        assert response.json()["project"]["id"] == project.id
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_report_project_context_accepts_workflow_only_permission():
    engine, db, client, project = _context()
    try:
        db.query(RolePermission).filter(
            RolePermission.permission_id == db.query(Permission.id).filter(
                Permission.code == "management_dashboard.view"
            ).scalar_subquery()
        ).delete(synchronize_session=False)
        db.commit()
        response = client.get("/reports/project-context", params={"project_id": project.id})
        assert response.status_code == 200
        assert response.json()["project"]["timezone"] == "Asia/Tehran"
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_report_project_context_denies_users_without_report_permission():
    engine, db, client, project = _context()
    try:
        db.query(RolePermission).delete()
        db.commit()
        response = client.get("/reports/project-context", params={"project_id": project.id})
        assert response.status_code == 403
        assert response.json()["detail"] == "forbidden"
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()



def test_management_summary_xlsx_export_is_valid_and_permission_scoped():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/management/summary/export.xlsx", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert 'filename="sabt-management-summary.xlsx"' in response.headers["content-disposition"]
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        assert sheet.sheet_view.rightToLeft is True
        assert [sheet.cell(row=1, column=i).value for i in (1, 2)] == ["شاخص", "مقدار"]
        labels = [sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)]
        assert "عملیات ایجادشده" in labels
        workbook.close()
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_workflow_xlsx_export_contains_persian_headers_and_is_project_scoped():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/workflow/export.xlsx", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 200
        assert response.content.startswith(b"PK")
        assert 'filename="sabt-workflow-report.xlsx"' in response.headers["content-disposition"]
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        assert sheet.cell(row=1, column=1).value == "شناسه عملیات"
        assert sheet.auto_filter.ref
        workbook.close()
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_workflow_events_xlsx_export_contains_event_headers():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/workflow/events/export.xlsx", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.active.cell(row=1, column=1).value == "شناسه رخداد"
        workbook.close()
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_xlsx_export_enforces_project_permissions():
    engine, db, client, project = _context()
    try:
        db.query(ProjectAccess).delete()
        db.commit()
        response = client.get("/reports/workflow/export.xlsx", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 403
        assert response.json()["detail"] == "forbidden"
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_xlsx_export_neutralizes_formula_cells():
    from app.reports.routes import _safe_xlsx_cell

    assert _safe_xlsx_cell("=1+1") == "'=1+1"
    assert _safe_xlsx_cell("ordinary text") == "ordinary text"
    assert _safe_xlsx_cell({"detail": "جزئیات"}) == '{"detail": "جزئیات"}'



def test_management_summary_xlsx_export_is_valid_and_permission_scoped():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/management/summary/export.xlsx", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert 'filename="sabt-management-summary.xlsx"' in response.headers["content-disposition"]
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        assert sheet.sheet_view.rightToLeft is True
        assert [sheet.cell(row=1, column=i).value for i in (1, 2)] == ["شاخص", "مقدار"]
        labels = [sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)]
        assert "عملیات ایجادشده" in labels
        workbook.close()
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_workflow_xlsx_export_contains_persian_headers_and_is_project_scoped():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/workflow/export.xlsx", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 200
        assert response.content.startswith(b"PK")
        assert 'filename="sabt-workflow-report.xlsx"' in response.headers["content-disposition"]
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        assert sheet.cell(row=1, column=1).value == "شناسه عملیات"
        assert sheet.auto_filter.ref
        workbook.close()
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_workflow_events_xlsx_export_contains_event_headers():
    engine, db, client, project = _context()
    try:
        response = client.get("/reports/workflow/events/export.xlsx", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.active.cell(row=1, column=1).value == "شناسه رخداد"
        workbook.close()
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_xlsx_export_enforces_project_permissions():
    engine, db, client, project = _context()
    try:
        db.query(ProjectAccess).delete()
        db.commit()
        response = client.get("/reports/workflow/export.xlsx", params={
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        })
        assert response.status_code == 403
        assert response.json()["detail"] == "forbidden"
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_xlsx_export_neutralizes_formula_cells():
    from app.reports.routes import _safe_xlsx_cell

    assert _safe_xlsx_cell("=1+1") == "'=1+1"
    assert _safe_xlsx_cell("ordinary text") == "ordinary text"
    assert _safe_xlsx_cell({"detail": "جزئیات"}) == '{"detail": "جزئیات"}'

def test_all_report_export_families_reject_users_without_project_access():
    engine, db, client, project = _context()
    try:
        db.query(ProjectAccess).delete()
        db.commit()
        common_params = {
            "project_id": project.id,
            "start_at": "2026-10-01T00:00:00+00:00",
            "end_at": "2026-11-01T00:00:00+00:00",
        }
        endpoints = [
            "/reports/management/summary/export.xlsx",
            "/reports/workflow/events/export.xlsx",
            "/reports/workflow/events/export.csv",
        ]
        for endpoint in endpoints:
            response = client.get(endpoint, params=common_params)
            assert response.status_code == 403, endpoint
            assert response.json()["detail"] == "forbidden", endpoint
    finally:
        app.dependency_overrides.clear()
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()
