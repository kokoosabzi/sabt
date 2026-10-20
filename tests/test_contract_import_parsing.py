import csv
import io

import pytest
from openpyxl import Workbook

from app.contracts.import_service import (
    ContractImportError,
    MAX_UPLOAD_BYTES,
    parse_parent_contract_import,
)


def test_parse_csv_normalizes_rows_and_optional_property():
    content = b"suffix,initial_person_id,property_id\n0123,7,4\n0124,8,\n"
    rows = parse_parent_contract_import("contracts.csv", content)
    assert rows == [
        {"row_number": 2, "suffix": "0123", "initial_person_id": 7, "property_id": 4},
        {"row_number": 3, "suffix": "0124", "initial_person_id": 8, "property_id": None},
    ]


def test_parse_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["suffix", "initial_person_id", "property_id"])
    sheet.append(["0100", 12, None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    rows = parse_parent_contract_import("contracts.xlsx", buffer.getvalue())
    assert rows[0]["suffix"] == "0100"
    assert rows[0]["initial_person_id"] == 12
    assert rows[0]["property_id"] is None


def test_parse_xlsx_preserves_leading_zeroes_from_four_digit_number_format():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["suffix", "initial_person_id", "property_id"])
    sheet.append([100, 12, None])
    sheet["A2"].number_format = "0000"
    buffer = io.BytesIO()
    workbook.save(buffer)

    rows = parse_parent_contract_import("contracts.xlsx", buffer.getvalue())

    assert rows[0]["suffix"] == "0100"
    assert rows[0]["initial_person_id"] == 12


def test_parse_xlsx_does_not_guess_missing_leading_zeroes_for_general_numeric_suffix():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["suffix", "initial_person_id"])
    sheet.append([100, 12])
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(ContractImportError, match="import_row_invalid:2"):
        parse_parent_contract_import("contracts.xlsx", buffer.getvalue())


@pytest.mark.parametrize(
    ("filename", "content", "error"),
    [
        ("contracts.txt", b"suffix,initial_person_id\n1,2", "import_file_type_unsupported"),
        ("contracts.csv", b"suffix\n1", "import_required_columns_missing"),
        ("contracts.csv", b"suffix,initial_person_id\n0001,abc", "import_row_invalid:2"),
        ("contracts.csv", b"suffix,initial_person_id\n", "import_no_data_rows"),
    ],
)
def test_rejects_invalid_import_files(filename, content, error):
    with pytest.raises(ContractImportError, match=error):
        parse_parent_contract_import(filename, content)


def test_rejects_oversized_upload():
    with pytest.raises(ContractImportError, match="import_file_too_large"):
        parse_parent_contract_import("contracts.csv", b"x" * (MAX_UPLOAD_BYTES + 1))


def test_rejects_invalid_csv_encoding():
    with pytest.raises(ContractImportError, match="import_csv_encoding_invalid"):
        parse_parent_contract_import("contracts.csv", b"\xff\xfe\xff")

@pytest.mark.parametrize(
    ("content", "error"),
    [
        (b"suffix,initial_person_id,initial_person_id\n0001,2,3\n", "import_duplicate_column_header"),
        (b"suffix,initial_person_id,\n0001,2,\n", "import_blank_column_header"),
        (b"suffix,initial_person_id\n0001,2,extra\n", "import_row_column_count_mismatch:2"),
    ],
)
def test_rejects_ambiguous_or_malformed_csv(content, error):
    with pytest.raises(ContractImportError, match=error):
        parse_parent_contract_import("contracts.csv", content)


def test_rejects_csv_over_row_limit_while_parsing():
    content = ("suffix,initial_person_id\n" + "".join(f"{i:04d},{i + 1}\n" for i in range(2001))).encode()
    with pytest.raises(ContractImportError, match="import_row_limit_exceeded"):
        parse_parent_contract_import("contracts.csv", content)


def test_rejects_xlsx_extension_with_non_zip_content():
    with pytest.raises(ContractImportError, match="import_file_invalid"):
        parse_parent_contract_import("contracts.xlsx", b"not a zip archive")

