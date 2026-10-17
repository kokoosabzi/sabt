from datetime import datetime
import csv
import io
import json

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.auth.dependencies import current_user, get_db
from app.reports.service import (
    ReportValidationError, get_management_summary, get_report_project_context,
    get_workflow_events_report, get_workflow_report,
)

router = APIRouter(prefix="/reports", tags=["reports"])
EXPORT_MAX_ROWS = 5000
EXPORT_PAGE_SIZE = 100


def _report_response(fn):
    try:
        return fn()
    except ReportValidationError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


def _safe_csv_cell(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        return value
    # Prevent spreadsheet formula execution for untrusted text fields.
    if value.lstrip(" \\t\\r\\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _csv_response(filename: str, columns: list[tuple[str, str]], rows: list[dict]) -> Response:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow([label for label, _ in columns])
    for row in rows:
        writer.writerow([_safe_csv_cell(row.get(key)) for _, key in columns])
    # UTF-8 BOM helps desktop Excel identify Persian text correctly.
    return Response(
        content="\ufeff" + buffer.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )



def _safe_xlsx_cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _xlsx_response(filename: str, columns: list[tuple[str, str]], rows: list[dict]) -> Response:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "گزارش"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.append([label for label, _ in columns])
    for row in rows:
        sheet.append([_safe_xlsx_cell(row.get(key)) for _, key in columns])
    if columns:
        sheet.auto_filter.ref = sheet.dimensions
        for index, (label, key) in enumerate(columns, start=1):
            max_length = len(label)
            for row in rows[:200]:
                value = _safe_xlsx_cell(row.get(key))
                max_length = max(max_length, len(str(value)) if value is not None else 0)
            sheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 42)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

def _collect_pages(fetch_page):
    first = fetch_page(0)
    total = first["pagination"]["total"]
    if total > EXPORT_MAX_ROWS:
        raise ReportValidationError("report_export_row_limit_exceeded")
    items = list(first["items"])
    for offset in range(EXPORT_PAGE_SIZE, total, EXPORT_PAGE_SIZE):
        items.extend(fetch_page(offset)["items"])
    return items


@router.get("/management/summary")
def management_summary(
    project_id: int = Query(..., ge=1),
    start_at: datetime = Query(...),
    end_at: datetime = Query(...),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    return _report_response(lambda: get_management_summary(db, user, project_id, start_at, end_at))


@router.get("/workflow")
def workflow_report(
    project_id: int = Query(..., ge=1),
    start_at: datetime = Query(...),
    end_at: datetime = Query(...),
    status: str | None = Query(None, min_length=1, max_length=32),
    operation_type_id: int | None = Query(None, ge=1),
    parent_contract_number: str | None = Query(None, min_length=6, max_length=6),
    child_contract_number: str | None = Query(None, min_length=8, max_length=32),
    property_id: int | None = Query(None, ge=1),
    created_by: int | None = Query(None, ge=1),
    limit: int = Query(50, ge=1, le=100),
    offset: int = Query(0, ge=0),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    return _report_response(lambda: get_workflow_report(
        db, user, project_id, start_at, end_at,
        status=status, operation_type_id=operation_type_id,
        parent_contract_number=parent_contract_number,
        child_contract_number=child_contract_number,
        property_id=property_id, created_by=created_by, limit=limit, offset=offset,
    ))


@router.get("/workflow/events")
def workflow_events_report(
    project_id: int = Query(..., ge=1),
    start_at: datetime = Query(...),
    end_at: datetime = Query(...),
    event_type: str | None = Query(None, min_length=1, max_length=64),
    operation_id: int | None = Query(None, ge=1),
    actor_user_id: int | None = Query(None, ge=1),
    limit: int = Query(50, ge=1, le=100),
    offset: int = Query(0, ge=0),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    return _report_response(lambda: get_workflow_events_report(
        db, user, project_id, start_at, end_at,
        event_type=event_type, operation_id=operation_id,
        actor_user_id=actor_user_id, limit=limit, offset=offset,
    ))


@router.get("/workflow/export.csv")
def export_workflow_csv(
    project_id: int = Query(..., ge=1),
    start_at: datetime = Query(...),
    end_at: datetime = Query(...),
    status: str | None = Query(None, min_length=1, max_length=32),
    operation_type_id: int | None = Query(None, ge=1),
    parent_contract_number: str | None = Query(None, min_length=6, max_length=6),
    child_contract_number: str | None = Query(None, min_length=8, max_length=32),
    property_id: int | None = Query(None, ge=1),
    created_by: int | None = Query(None, ge=1),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    def build():
        items = _collect_pages(lambda offset: get_workflow_report(
            db, user, project_id, start_at, end_at,
            status=status, operation_type_id=operation_type_id,
            parent_contract_number=parent_contract_number,
            child_contract_number=child_contract_number,
            property_id=property_id, created_by=created_by,
            limit=EXPORT_PAGE_SIZE, offset=offset,
        ))
        columns = [
            ("شناسه عملیات", "operation_id"), ("نوع عملیات", "operation_type_name"),
            ("وضعیت", "status"), ("وضعیت گردش کار", "workflow_state"),
            ("شناسه ملک", "property_id"), ("شماره قرارداد مادر", "parent_contract_number"),
            ("شماره انتقال", "child_contract_number"), ("ثبت‌کننده", "created_by_display_name"),
            ("تاریخ ایجاد", "created_at"), ("تاریخ به‌روزرسانی", "updated_at"),
            ("تاریخ نهایی‌شدن", "finalized_at"), ("شماره نسخه", "current_version_number"),
        ]
        return _csv_response("sabt-workflow-report.csv", columns, items)
    return _report_response(build)


@router.get("/workflow/events/export.csv")
def export_workflow_events_csv(
    project_id: int = Query(..., ge=1),
    start_at: datetime = Query(...),
    end_at: datetime = Query(...),
    event_type: str | None = Query(None, min_length=1, max_length=64),
    operation_id: int | None = Query(None, ge=1),
    actor_user_id: int | None = Query(None, ge=1),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    def build():
        items = _collect_pages(lambda offset: get_workflow_events_report(
            db, user, project_id, start_at, end_at,
            event_type=event_type, operation_id=operation_id,
            actor_user_id=actor_user_id, limit=EXPORT_PAGE_SIZE, offset=offset,
        ))
        columns = [
            ("شناسه رخداد", "event_id"), ("نوع رخداد", "event_type"),
            ("شناسه عملیات", "operation_id"), ("نوع عملیات", "operation_type_name"),
            ("شماره قرارداد مادر", "parent_contract_number"), ("شماره انتقال", "child_contract_number"),
            ("شناسه کاربر", "actor_user_id"), ("نام کاربر", "actor_display_name"),
            ("وضعیت قبل", "from_status"), ("وضعیت بعد", "to_status"),
            ("زمان رخداد", "occurred_at"), ("جزئیات", "metadata"),
        ]
        for item in items:
            if isinstance(item.get("metadata"), dict):
                item["metadata"] = str(item["metadata"])
        return _csv_response("sabt-workflow-events.csv", columns, items)
    return _report_response(build)


@router.get("/management/summary/export.xlsx")
def export_management_summary_xlsx(
    project_id: int = Query(..., ge=1),
    start_at: datetime = Query(...),
    end_at: datetime = Query(...),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    def build():
        summary = get_management_summary(db, user, project_id, start_at, end_at)
        indicators = summary["indicators"]
        rows = [
            {"label": "شناسه پروژه", "value": summary["project"]["id"]},
            {"label": "کد پروژه", "value": summary["project"]["code"]},
            {"label": "نام پروژه", "value": summary["project"]["name"]},
            {"label": "منطقه زمانی", "value": summary["project"]["timezone"]},
            {"label": "شروع بازه (UTC)", "value": summary["period"]["start_at"]},
            {"label": "پایان بازه (UTC)", "value": summary["period"]["end_at"]},
            {"label": "عملیات ایجادشده", "value": indicators["created_operations"]},
            {"label": "عملیات نهایی‌شده", "value": indicators["finalized_operations"]},
            {"label": "عملیات باز فعلی", "value": indicators["currently_open_operations"]},
            {"label": "قراردادهای مادر ایجادشده", "value": indicators["parent_contracts_created"]},
            {"label": "انتقال‌های نهایی‌شده", "value": indicators["transfers_finalized"]},
        ]
        rows.extend(
            {"label": "عملیات ایجادشده - وضعیت " + item["status"], "value": item["count"]}
            for item in indicators["created_operations_by_status"]
        )
        return _xlsx_response("sabt-management-summary.xlsx", [("شاخص", "label"), ("مقدار", "value")], rows)
    return _report_response(build)


@router.get("/workflow/export.xlsx")
def export_workflow_xlsx(
    project_id: int = Query(..., ge=1),
    start_at: datetime = Query(...),
    end_at: datetime = Query(...),
    status: str | None = Query(None, min_length=1, max_length=32),
    operation_type_id: int | None = Query(None, ge=1),
    parent_contract_number: str | None = Query(None, min_length=6, max_length=6),
    child_contract_number: str | None = Query(None, min_length=8, max_length=32),
    property_id: int | None = Query(None, ge=1),
    created_by: int | None = Query(None, ge=1),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    def build():
        items = _collect_pages(lambda offset: get_workflow_report(
            db, user, project_id, start_at, end_at, status=status,
            operation_type_id=operation_type_id, parent_contract_number=parent_contract_number,
            child_contract_number=child_contract_number, property_id=property_id,
            created_by=created_by, limit=EXPORT_PAGE_SIZE, offset=offset,
        ))
        columns = [
            ("شناسه عملیات", "operation_id"), ("نوع عملیات", "operation_type_name"),
            ("وضعیت", "status"), ("وضعیت گردش کار", "workflow_state"),
            ("شناسه ملک", "property_id"), ("شماره قرارداد مادر", "parent_contract_number"),
            ("شماره انتقال", "child_contract_number"), ("ثبت‌کننده", "created_by_display_name"),
            ("تاریخ ایجاد", "created_at"), ("تاریخ به‌روزرسانی", "updated_at"),
            ("تاریخ نهایی‌شدن", "finalized_at"), ("شماره نسخه", "current_version_number"),
        ]
        return _xlsx_response("sabt-workflow-report.xlsx", columns, items)
    return _report_response(build)


@router.get("/workflow/events/export.xlsx")
def export_workflow_events_xlsx(
    project_id: int = Query(..., ge=1),
    start_at: datetime = Query(...),
    end_at: datetime = Query(...),
    event_type: str | None = Query(None, min_length=1, max_length=64),
    operation_id: int | None = Query(None, ge=1),
    actor_user_id: int | None = Query(None, ge=1),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    def build():
        items = _collect_pages(lambda offset: get_workflow_events_report(
            db, user, project_id, start_at, end_at, event_type=event_type,
            operation_id=operation_id, actor_user_id=actor_user_id,
            limit=EXPORT_PAGE_SIZE, offset=offset,
        ))
        columns = [
            ("شناسه رخداد", "event_id"), ("نوع رخداد", "event_type"),
            ("شناسه عملیات", "operation_id"), ("نوع عملیات", "operation_type_name"),
            ("شماره قرارداد مادر", "parent_contract_number"), ("شماره انتقال", "child_contract_number"),
            ("شناسه کاربر", "actor_user_id"), ("نام کاربر", "actor_display_name"),
            ("وضعیت قبل", "from_status"), ("وضعیت بعد", "to_status"),
            ("زمان رخداد", "occurred_at"), ("جزئیات", "metadata"),
        ]
        return _xlsx_response("sabt-workflow-events.xlsx", columns, items)
    return _report_response(build)


@router.get("/project-context")
def report_project_context(
    project_id: int = Query(..., ge=1),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    return _report_response(lambda: get_report_project_context(db, user, project_id))
