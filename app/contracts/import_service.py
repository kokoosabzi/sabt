"""Safe CSV/XLSX staging for parent-contract bulk imports.

Input columns: suffix, initial_person_id, optional property_id.
Parsing and preview are side-effect free. Confirmation must revalidate in the
same transaction used to create every parent contract.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import Counter
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth.authorization import AuthorizationContext, authorize
from app.contracts.numbering import (
    build_parent_contract_number,
    validate_parent_contract_suffix,
    validate_project_contract_code,
)
from app.db.models import ParentContract, Person, Project, ProjectPerson, ProjectProperty, Property, User


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ZIP_ENTRIES = 200
REQUIRED_COLUMNS = {"suffix", "initial_person_id"}
ALLOWED_COLUMNS = REQUIRED_COLUMNS | {"property_id"}


class ContractImportError(ValueError):
    """An import file or row cannot be safely processed."""


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_parent_contract_import(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Parse CSV or XLSX bytes into normalized rows without database writes."""
    if not content:
        raise ContractImportError("import_file_empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ContractImportError("import_file_too_large")
    name = (filename or "").lower()
    try:
        if name.endswith(".csv"):
            text = content.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text))
            raw_headers = next(reader, None)
            if not raw_headers:
                raise ContractImportError("import_required_columns_missing")
            headers = [str(header).strip().lower() for header in raw_headers]
            if any(not header for header in headers):
                raise ContractImportError("import_blank_column_header")
            if len(headers) != len(set(headers)):
                raise ContractImportError("import_duplicate_column_header")
            header_set = set(headers)
            if not REQUIRED_COLUMNS.issubset(header_set):
                raise ContractImportError("import_required_columns_missing")
            if header_set - ALLOWED_COLUMNS:
                raise ContractImportError("import_unknown_columns")
            rows = []
            for row_number, values in enumerate(reader, start=2):
                if len(values) != len(headers):
                    raise ContractImportError(f"import_row_column_count_mismatch:{row_number}")
                rows.append(dict(zip(headers, values)))
                if len(rows) > MAX_IMPORT_ROWS:
                    raise ContractImportError("import_row_limit_exceeded")
        elif name.endswith(".xlsx"):
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    entries = archive.infolist()
                    if len(entries) > MAX_XLSX_ZIP_ENTRIES:
                        raise ContractImportError("import_xlsx_too_many_zip_entries")
                    if any(entry.flag_bits & 0x1 for entry in entries):
                        raise ContractImportError("import_xlsx_encrypted")
                    if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                        raise ContractImportError("import_xlsx_expanded_size_too_large")
                    if any(entry.file_size > 1024 * 1024 and entry.compress_size == 0 for entry in entries):
                        raise ContractImportError("import_xlsx_suspicious_compression")
            except zipfile.BadZipFile as exc:
                raise ContractImportError("import_file_invalid") from exc
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                iterator = sheet.iter_rows()
                first = next(iterator, None)
                headers = [
                    str(cell.value).strip().lower() if cell.value is not None else ""
                    for cell in (first or ())
                ]
                if any(not header for header in headers):
                    raise ContractImportError("import_blank_column_header")
                if len(headers) != len(set(headers)):
                    raise ContractImportError("import_duplicate_column_header")
                header_set = set(headers)
                if not REQUIRED_COLUMNS.issubset(header_set):
                    raise ContractImportError("import_required_columns_missing")
                if header_set - ALLOWED_COLUMNS:
                    raise ContractImportError("import_unknown_columns")
                rows = []
                for row_number, cells in enumerate(iterator, start=2):
                    if all(cell.value is None for cell in cells):
                        continue
                    if len(cells) != len(headers):
                        raise ContractImportError(f"import_row_column_count_mismatch:{row_number}")
                    values: dict[str, Any] = {}
                    for header, cell in zip(headers, cells):
                        value = cell.value
                        # Excel stores a numeric 100 as 100 even when displayed as 0100.
                        # Honor an explicit four-zero display mask for the four-digit
                        # business suffix, without guessing leading zeroes for General cells.
                        if (
                            header == "suffix"
                            and isinstance(value, (int, float))
                            and not isinstance(value, bool)
                            and float(value).is_integer()
                            and re.fullmatch(r"0{4}", cell.number_format or "") is not None
                        ):
                            value = f"{int(value):04d}"
                        values[header] = value
                    rows.append(values)
                    if len(rows) > MAX_IMPORT_ROWS:
                        raise ContractImportError("import_row_limit_exceeded")
            finally:
                workbook.close()
        else:
            raise ContractImportError("import_file_type_unsupported")
    except UnicodeDecodeError as exc:
        raise ContractImportError("import_csv_encoding_invalid") from exc
    except ContractImportError:
        raise
    except Exception as exc:
        raise ContractImportError("import_file_invalid") from exc

    if len(rows) > MAX_IMPORT_ROWS:
        raise ContractImportError("import_row_limit_exceeded")
    normalized: list[dict[str, Any]] = []
    for index, row in enumerate(rows, start=2):
        suffix = _text(row.get("suffix"))
        person_raw = _text(row.get("initial_person_id"))
        property_raw = _text(row.get("property_id"))
        try:
            suffix = validate_parent_contract_suffix(suffix)
            person_id = int(person_raw)
            property_id = int(property_raw) if property_raw else None
            if person_id <= 0 or (property_id is not None and property_id <= 0):
                raise ValueError
        except (ValueError, TypeError) as exc:
            raise ContractImportError(f"import_row_invalid:{index}") from exc
        normalized.append({
            "row_number": index,
            "suffix": suffix,
            "initial_person_id": person_id,
            "property_id": property_id,
        })
    if not normalized:
        raise ContractImportError("import_no_data_rows")
    return normalized


def preview_parent_contract_import(
    db: Session, user: User, project_id: int, rows: list[dict[str, Any]]
) -> dict[str, Any]:
    """Validate import rows and report duplicates without mutating the database."""
    authorize(db, user, "parent_contract.create", AuthorizationContext(project_id=project_id))
    project = db.get(Project, project_id)
    if project is None:
        raise ContractImportError("project_not_found")
    if project.status != "ACTIVE":
        raise ContractImportError("project_inactive")
    if project.contract_code is None:
        raise ContractImportError("project_contract_code_required")
    code = validate_project_contract_code(project.contract_code)

    suffix_counts = Counter(row["suffix"] for row in rows)
    issues: list[dict[str, Any]] = []
    preview: list[dict[str, Any]] = []
    for row in rows:
        row_issues: list[str] = []
        suffix = row["suffix"]
        number = build_parent_contract_number(code, suffix)
        if suffix_counts[suffix] > 1:
            row_issues.append("duplicate_suffix_in_file")
        if db.scalar(select(ParentContract.id).where(
            ParentContract.project_id == project_id, ParentContract.suffix == suffix
        )) is not None:
            row_issues.append("suffix_already_exists")
        if db.scalar(select(ParentContract.id).where(
            ParentContract.contract_number == number
        )) is not None:
            row_issues.append("contract_number_already_exists")
        person = db.get(Person, row["initial_person_id"])
        membership = db.get(ProjectPerson, (project_id, row["initial_person_id"]))
        if person is None or not person.is_active:
            row_issues.append("initial_person_invalid")
        elif membership is None or not membership.active:
            row_issues.append("initial_person_not_associated_with_project")
        property_id = row["property_id"]
        if property_id is not None:
            prop = db.get(Property, property_id)
            association = db.get(ProjectProperty, (project_id, property_id))
            if prop is None or not prop.is_active:
                row_issues.append("property_invalid")
            elif association is None or not association.active:
                row_issues.append("property_not_associated_with_project")
        item = {**row, "contract_number": number, "issues": row_issues, "valid": not row_issues}
        preview.append(item)
        for issue in row_issues:
            issues.append({"row_number": row["row_number"], "code": issue})
    return {
        "project_id": project_id,
        "row_count": len(rows),
        "valid_count": sum(1 for row in preview if row["valid"]),
        "invalid_count": sum(1 for row in preview if not row["valid"]),
        "can_confirm": bool(preview) and not issues,
        "rows": preview,
        "issues": issues,
    }
